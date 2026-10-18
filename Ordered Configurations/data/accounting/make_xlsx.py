"""make_xlsx.py -- builds accounting.xlsx from accounting.csv, accounting_graphs.csv and (optionally)
reverification.csv.  The Summary sheet counts rows of the 'All computations' sheet with COUNTIFS
formulas, so the counts update if rows are edited."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
F = Font(name="Arial", size=10); B = Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", fgColor="DDE3EA")

def sheet(wb, title, path, widths):
    ws = wb.create_sheet(title)
    rows = list(csv.reader(open(path)))
    for r in rows:
        ws.append([int(x) if x.isdigit() else x for x in r])
    for c in ws[1]:
        c.font = B; c.fill = HEAD
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws, len(rows) - 1

wb = Workbook(); wb.remove(wb.active)
summ = wb.create_sheet("Summary")
ws, N = sheet(wb, "All computations", os.path.join(HERE, "accounting.csv"), [24, 30, 30, 5, 5, 7, 4, 60, 10, 50, 22, 36])
gs, G = sheet(wb, "Distinct graphs", os.path.join(HERE, "accounting_graphs.csv"), [30, 24, 40, 5, 5, 7, 12, 16, 80])
if os.path.exists(os.path.join(HERE, "reverification.csv")):
    sheet(wb, "Re-verification Oct 2026", os.path.join(HERE, "reverification.csv"), [36, 8, 50, 50, 14, 10])
rng = lambda col: f"'All computations'!${col}$2:${col}${N + 1}"
summ["A1"] = "Accounting of all recorded computations"; summ["A1"].font = Font(name="Arial", size=12, bold=True)
summ["A2"] = "Generated from data/accounting/accounting.csv by make_xlsx.py; counts below are formulas over the 'All computations' sheet."
summ["A2"].font = F
r = 4
for label, f in [("Computations recorded", f"=COUNTA({rng('A')})"),
                 ("  complete", f"=COUNTIF({rng('K')},\"complete\")"),
                 ("  failed (out of memory)", f"=COUNTIF({rng('K')},\"failed (memory)\")"),
                 ("  partial", f"=COUNTIF({rng('K')},\"partial*\")"),
                 ("  not completed (cancelled)", f"=COUNTIF({rng('K')},\"not completed\")"),
                 ("Distinct graphs up to isomorphism", f"=COUNTA('Distinct graphs'!$A$2:$A${G + 1})"),
                 ("  planar", f"=COUNTIF('Distinct graphs'!$F$2:$F${G + 1},\"yes\")"),
                 ("  non-planar", f"=COUNTIF('Distinct graphs'!$F$2:$F${G + 1},\"no\")"),
                 ("Computations on non-planar graphs", f"=COUNTIF({rng('F')},\"no\")")]:
    summ.cell(r, 1, label).font = F; summ.cell(r, 2, f).font = F; r += 1
r += 1
summ.cell(r, 1, "By category").font = B; summ.cell(r, 2, "computations").font = B; summ.cell(r, 3, "complete").font = B; r += 1
cats = []
for row in csv.DictReader(open(os.path.join(HERE, "accounting.csv"))):
    if row["category"] not in cats: cats.append(row["category"])
for c in cats:
    summ.cell(r, 1, c).font = F
    summ.cell(r, 2, f"=COUNTIF({rng('A')},A{r})").font = F
    summ.cell(r, 3, f"=COUNTIFS({rng('A')},A{r},{rng('K')},\"complete\")").font = F
    r += 1
r += 1
summ.cell(r, 1, "By number of particles n").font = B; r += 1
for n in range(2, 9):
    summ.cell(r, 1, n).font = F; summ.cell(r, 2, f"=COUNTIF({rng('G')},A{r})").font = F; r += 1
summ.column_dimensions["A"].width = 40; summ.column_dimensions["B"].width = 14; summ.column_dimensions["C"].width = 12
wb.save(os.path.join(HERE, "accounting.xlsx"))
print("rows", N, "graphs", G)
